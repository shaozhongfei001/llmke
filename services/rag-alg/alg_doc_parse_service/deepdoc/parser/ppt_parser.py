from io import BytesIO
from typing import Iterator

from langchain_core.document_loaders import Blob
from pptx import Presentation

from .base_parser import BaseParser
from ..output_object import ParsedTextMetadata, ParsedText


class _BaseParser(BaseParser):
    def lazy_parse(self, blob: Blob) -> Iterator[ParsedText]:
        slides = self(blob)

        yield from [
            ParsedText(page_content=slide,
                       metadata=ParsedTextMetadata(source=blob.source, file_path=blob.source)
                       ) for slide in slides
        ]

    def _extract(self, shape):
        if shape.shape_type == 19:
            tb = shape.table
            rows = []
            for i in range(1, len(tb.rows)):
                rows.append("; ".join([tb.cell(
                    0, j).text + ": " + tb.cell(i, j).text for j in range(len(tb.columns)) if tb.cell(i, j)]))
            return "\n".join(rows)

        if shape.has_text_frame:
            return shape.text_frame.text

        if shape.shape_type == 6:
            texts = []
            for p in sorted(shape.shapes, key=lambda x: (x.top // 10, x.left)):
                t = self._extract(p)
                if t:
                    texts.append(t)
            return "\n".join(texts)

    def __call__(self, file: str | Blob):
        ppt = Presentation(file) if isinstance(file, str) else Presentation(BytesIO(file.as_bytes()))
        slide_texts = []
        for i, slide in enumerate(ppt.slides):
            text_parts = []
            for shape in sorted(slide.shapes, key=lambda x: (x.top // 10, x.left) if x.top else (0, x.left)):
                text_part = self._extract(shape)
                if text_part:
                    text_parts.append(text_part)
            slide_texts.append('\n'.join(text_parts))

        return slide_texts


class Parser:
    @staticmethod
    def get_parser() -> _BaseParser:
        return _BaseParser()
