import openpyxl
from openpyxl.utils import get_column_letter

file_path = "test_exl.xlsx"


def read_excel(file_path):
    workbook = openpyxl.load_workbook(file_path)

    sheet_names = workbook.sheetnames
    all_blocks = []
    for sheet_name in sheet_names:
        print(sheet_name)
        sheet = workbook[sheet_name]
        max_row = sheet.max_row
        max_column = sheet.max_column

        visited = [[False for _ in range(max_column)] for _ in range(max_row)]
        merged_cells = list(sheet.merged_cells.ranges)

        for r in range(1, max_row + 1):
            for c in range(1, max_column + 1):
                if not visited[r - 1][c - 1]:
                    current_block = []
                    dfs(r, c, current_block, sheet, max_row, max_column, visited, merged_cells)
                    if current_block:
                        all_blocks.append(current_block)
    print(len(all_blocks))
    for block in all_blocks:
        print(block)

    return all_blocks


def dfs(row, col, current_block, sheet, max_row, max_column, visited, merged_cells):
    if row < 1 or row > max_row or col < 1 or col > max_column or visited[row - 1][col - 1]:
        return

    visited[row - 1][col - 1] = True

    # 检查是否为合并单元格
    merged_range = get_merged_range(row, col, merged_cells)
    if merged_range:
        min_row, min_col, max_row, max_col = merged_range.min_row, merged_range.min_col, merged_range.max_row, merged_range.max_col

        cell_value = sheet.cell(min_row, min_col).value

        if cell_value is None:
            return

        cell_raw_column = []
        for r in range(min_row, max_row + 1):
            for c in range(min_col, max_col + 1):
                visited[r - 1][c - 1] = True
                cell_raw_column.append((r, c))

        # 合并区域的偏移
        cell_offset = f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}"

        current_block.append({
            'cell_offset': cell_offset,
            'cell_values': cell_value,
            'cell_raw_column': cell_raw_column
        })

        return

    # 如果不是合并单元格，则处理单个单元格
    cell_value = sheet.cell(row, col).value

    if cell_value is None:
        return

    # 单元格的偏移
    cell_offset = f"{get_column_letter(col)}{row}"

    current_block.append({
        'cell_offset': cell_offset,
        'cell_values': cell_value,
        'cell_raw_column': [(row, col)]
    })

    dfs(row + 1, col, current_block, sheet, max_row, max_column, visited, merged_cells)
    dfs(row - 1, col, current_block, sheet, max_row, max_column, visited, merged_cells)
    dfs(row, col + 1, current_block, sheet, max_row, max_column, visited, merged_cells)
    dfs(row, col - 1, current_block, sheet, max_row, max_column, visited, merged_cells)


# 判断当前单元格是不是合并的
def get_merged_range(row, col, merged_cells):
    cell_address = f"{get_column_letter(col)}{row}"
    for merged_range in merged_cells:
        if cell_address in merged_range:
            return merged_range
    return None


if __name__ == '__main__':
    read_excel(file_path)
