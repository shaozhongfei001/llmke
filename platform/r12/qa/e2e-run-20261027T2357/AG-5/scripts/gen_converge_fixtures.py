#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""S14 收敛夹具：同名文件的两个版本。
  v0/R12FIX-product-loan.docx  原件（V0 内容：最高额度 5000万元）
  v1/R12FIX-product-loan.docx  改写件（V1 内容：最高额度 9000万元 + 唯一标记 AG5-CONV-MARK-9001）

要点：文件名必须相同（版本归组靠 parent_file_name），内容必须有可区分的唯一串，
否则"检索命中新版本"无法判定（只判 200 会得到假绿）。
"""
import os
import shutil
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from ag5_lib import EVID, FIXTURES, dump, sha256

BASE = os.path.join(EVID, "fixtures", "converge")
V0 = os.path.join(BASE, "v0")
V1 = os.path.join(BASE, "v1")
for d in (V0, V1):
    os.makedirs(d, exist_ok=True)

NAME = "R12FIX-product-loan.docx"
shutil.copy2(FIXTURES["docx_loan"], os.path.join(V0, NAME))

from docx import Document  # noqa: E402

MARK = "AG5-CONV-MARK-9001"
doc = Document()
doc.add_heading("智造贷 产品要点（AG5 收敛验证 V1）", level=1)
doc.add_paragraph("本文件用于验证同名新版本上传后的收敛行为，唯一标记：" + MARK)
doc.add_paragraph("适用对象：制造业中小微企业")
doc.add_paragraph("最高额度：9000万元（信用部分不超过1000万元）")
doc.add_paragraph("期限：最长3年")
doc.add_paragraph("利率：LPR+50BP 起")
doc.save(os.path.join(V1, NAME))

res = {
    "marker": MARK,
    "v0": {"path": os.path.join(V0, NAME), "bytes": os.path.getsize(os.path.join(V0, NAME)),
           "sha256": sha256(os.path.join(V0, NAME)), "content_hint": "原件：最高额度 5000万元"},
    "v1": {"path": os.path.join(V1, NAME), "bytes": os.path.getsize(os.path.join(V1, NAME)),
           "sha256": sha256(os.path.join(V1, NAME)),
           "content_hint": f"改写件：最高额度 9000万元 + 标记 {MARK}"},
}
dump("converge-fixtures.json", res, sub="S14")
print(res["v0"]["path"], res["v0"]["sha256"])
print(res["v1"]["path"], res["v1"]["sha256"])
